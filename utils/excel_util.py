from openpyxl import load_workbook

def open_excel_file(file_path,sheet_name):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]
    data = []

    for row in range(2,sheet.max_row+1):
        username = sheet.cell(row,column=1).value
        password = sheet.cell(row,column=2).value
        expected = sheet.cell(row,column=3).value
        data.append([username,password,expected])
    return data

def write_test_result(file_path,sheet_name,test_case_id,test_result):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    for row in range(2,sheet.max_row+1):
        if sheet.cell(row=row,column=1).value==test_case_id:
            sheet.cell(row=row,column=9).value=test_result
            break
    workbook.save(file_path)